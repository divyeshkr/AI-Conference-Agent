"""Verify every file-type parser produces artifacts. Run after changing core/ingest.py."""

from __future__ import annotations

import io

from core.ingest import ingest


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for title, body in [
        ("TRIUMPH-4 Phase 3 Results", "Mean weight reduction -24.2% at week 48"),
        ("Safety Profile", "GI adverse events 42%; discontinuation 6.1%"),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = title
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
        box.text_frame.text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Session summary: oral GLP-1 landscape", fontsize=14)
    page.insert_text((72, 130), "Aleniglipron -12.1% placebo-adjusted at week 36", fontsize=11)
    page.insert_text((72, 160), "Tolerability remains the differentiating factor", fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("1:1 with Amgen medical affairs, booth 412.")
    doc.add_paragraph("MariTide monthly dosing is the whole pitch this year.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_png() -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1200, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((60, 60), "POSTER 412: CagriSema cardiometabolic outcomes", fill="black")
    draw.text((60, 120), "-13.7% body weight | HbA1c -2.2pp | SBP -7.4 mmHg", fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


CASES = [
    ("deck.pptx", make_pptx),
    ("session.pdf", make_pdf),
    ("notes.docx", make_docx),
    ("poster.png", make_png),
    ("quick.txt", lambda: b"Payers want outcomes, not weight loss."),
]


def main() -> None:
    for name, factory in CASES:
        artifacts = ingest(name, factory())
        assert artifacts, f"{name} produced no artifacts"
        first = artifacts[0]
        detail = (
            f"{len(first.image_bytes)} image bytes"
            if first.image_bytes
            else f"{len(first.text)} chars: {first.text[:48]!r}"
        )
        print(f"{name:<14} -> {len(artifacts)} artifact(s) | {first.source_ref or '-':<8} | {detail}")
    print("\nALL PARSERS OK")


if __name__ == "__main__":
    main()
