"""Turn uploaded files into normalised Artifacts ready for extraction.

Photos keep their bytes so a vision model can read poster layouts directly;
classic OCR mangles multi-column posters and chart labels.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from .config import THUMB_DIR, UPLOAD_DIR
from .schema import SourceType

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".heic"}
TEXT_EXT = {".txt", ".md", ".csv"}


@dataclass
class Artifact:
    """One extractable unit: a photo, a single slide, a PDF page, a note."""

    source_file: str
    source_type: SourceType
    source_ref: str = ""
    text: str = ""
    image_bytes: bytes | None = None
    thumbnail_path: str | None = None
    meta: dict = field(default_factory=dict)

    def label(self) -> str:
        return f"{self.source_file}{f' · {self.source_ref}' if self.source_ref else ''}"


def save_upload(name: str, data: bytes) -> Path:
    dest = UPLOAD_DIR / name
    dest.write_bytes(data)
    return dest


def _thumbnail(data: bytes, stem: str) -> str | None:
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(data))
        img.thumbnail((480, 480))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        out = THUMB_DIR / f"{stem}.jpg"
        img.save(out, "JPEG", quality=82)
        return str(out)
    except Exception:  # noqa: BLE001 - thumbnails are cosmetic
        return None


def _compress_image(data: bytes, max_side: int = 1600) -> bytes:
    """Keep vision-model payloads small without losing poster legibility."""
    try:
        from PIL import Image

        img = Image.open(io.BytesIO(data))
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        if max(img.size) > max_side:
            ratio = max_side / max(img.size)
            img = img.resize((int(img.width * ratio), int(img.height * ratio)))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=85)
        return buf.getvalue()
    except Exception:  # noqa: BLE001
        return data


def _from_image(name: str, data: bytes) -> list[Artifact]:
    stem = Path(name).stem
    return [
        Artifact(
            source_file=name,
            source_type=SourceType.PHOTO,
            source_ref="photo",
            image_bytes=_compress_image(data),
            thumbnail_path=_thumbnail(data, stem),
        )
    ]


def _from_pptx(name: str, data: bytes) -> list[Artifact]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out: list[Artifact] = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    chunks.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            chunks.append("SPEAKER NOTES: " + slide.notes_slide.notes_text_frame.text.strip())
        text = "\n".join(chunks).strip()
        if text:
            out.append(
                Artifact(
                    source_file=name,
                    source_type=SourceType.SLIDE,
                    source_ref=f"slide {i}",
                    text=text,
                )
            )
    return out


def _from_pdf(name: str, data: bytes) -> list[Artifact]:
    import pymupdf

    doc = pymupdf.open(stream=data, filetype="pdf")
    stem = Path(name).stem
    out: list[Artifact] = []
    for i, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if len(text) >= 60:
            out.append(
                Artifact(
                    source_file=name,
                    source_type=SourceType.PDF_PAGE,
                    source_ref=f"page {i}",
                    text=text,
                )
            )
        else:
            # Scanned or image-only page: rasterise and let the vision model read it.
            png = page.get_pixmap(dpi=150).tobytes("png")
            out.append(
                Artifact(
                    source_file=name,
                    source_type=SourceType.PDF_PAGE,
                    source_ref=f"page {i}",
                    image_bytes=_compress_image(png),
                    thumbnail_path=_thumbnail(png, f"{stem}_p{i}"),
                )
            )
    doc.close()
    return out


def _from_docx(name: str, data: bytes) -> list[Artifact]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    chunks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(chunks).strip()
    return _chunk_note(name, text)


def _chunk_note(name: str, text: str, max_chars: int = 4000) -> list[Artifact]:
    if not text:
        return []
    if len(text) <= max_chars:
        return [Artifact(source_file=name, source_type=SourceType.NOTE, text=text)]
    out: list[Artifact] = []
    paragraphs = text.split("\n")
    buf: list[str] = []
    size = 0
    part = 1
    for para in paragraphs:
        if size + len(para) > max_chars and buf:
            out.append(
                Artifact(
                    source_file=name,
                    source_type=SourceType.NOTE,
                    source_ref=f"part {part}",
                    text="\n".join(buf),
                )
            )
            buf, size, part = [], 0, part + 1
        buf.append(para)
        size += len(para)
    if buf:
        out.append(
            Artifact(
                source_file=name,
                source_type=SourceType.NOTE,
                source_ref=f"part {part}",
                text="\n".join(buf),
            )
        )
    return out


def ingest(name: str, data: bytes) -> list[Artifact]:
    """Route an uploaded file to the right parser."""
    ext = Path(name).suffix.lower()
    if ext in IMAGE_EXT:
        return _from_image(name, data)
    if ext == ".pptx":
        return _from_pptx(name, data)
    if ext == ".pdf":
        return _from_pdf(name, data)
    if ext == ".docx":
        return _from_docx(name, data)
    if ext in TEXT_EXT:
        return _chunk_note(name, data.decode("utf-8", errors="replace"))
    raise ValueError(f"Unsupported file type: {ext}")


SUPPORTED_EXTENSIONS = sorted(
    e.lstrip(".") for e in IMAGE_EXT | TEXT_EXT | {".pptx", ".pdf", ".docx"}
)
